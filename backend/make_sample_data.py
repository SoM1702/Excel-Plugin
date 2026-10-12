"""Generate a realistic, explainable customer-churn demo dataset.

The labels are simulated from nine customer-behaviour signals. They are not
copied from a real company, but the relationships are deliberately plausible.

Run: ``python make_sample_data.py``
Outputs: backend/sample_data.csv and ../sample_data.xlsx.
"""

from pathlib import Path

import numpy as np
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill


rng = np.random.default_rng(42)
n = 400

# All features are numeric so the add-in can feed them directly into logistic
# regression without a categorical-encoding step.
age = rng.normal(40, 12, n).clip(18, 80)
monthly_charge = rng.normal(70, 25, n).clip(15, 150)
tenure_months = rng.normal(24, 18, n).clip(0, 72)
support_calls = rng.poisson(2, n).clip(0, 12)
late_payments = rng.poisson(0.8, n).clip(0, 6)
usage_drop_pct = (rng.beta(2.2, 4.5, n) * 100).clip(0, 100)
complaints_last_3_months = rng.poisson(0.7, n).clip(0, 6)
monthly_contract = rng.binomial(1, 0.55, n)  # 1=monthly rolling, 0=annual/fixed
days_since_last_login = rng.gamma(2.0, 5.5, n).clip(0, 60)


def zscore(values: np.ndarray) -> np.ndarray:
    return (values - values.mean()) / (values.std() + 1e-9)


# Simulated business rule used only to create a realistic, learnable outcome.
# It retains randomness, so the data is not perfectly separable or misleadingly
# easy. The add-in's held-out score is roughly 80% with default settings.
signal_strength = 1.55
z = -0.35 + signal_strength * (
    0.18 * zscore(age)
    + 0.38 * zscore(monthly_charge)
    - 0.52 * zscore(tenure_months)
    + 0.30 * zscore(support_calls)
    + 0.70 * zscore(late_payments)
    + 0.92 * zscore(usage_drop_pct)
    + 0.63 * zscore(complaints_last_3_months)
    + 0.60 * monthly_contract
    + 0.72 * zscore(days_since_last_login)
)
churn_probability = 1 / (1 + np.exp(-z))
churn = (rng.random(n) < churn_probability).astype(int)

headers = [
    "age", "monthly_charge", "tenure_months", "support_calls",
    "late_payments", "usage_drop_pct", "complaints_last_3_months",
    "monthly_contract", "days_since_last_login", "churn",
]
data = np.column_stack([
    age.round(0), monthly_charge.round(2), tenure_months.round(0),
    support_calls, late_payments, usage_drop_pct.round(1),
    complaints_last_3_months, monthly_contract, days_since_last_login.round(0),
    churn,
])

root = Path(__file__).resolve().parent.parent
csv_path = Path(__file__).with_name("sample_data.csv")
xlsx_path = root / "sample_data.xlsx"

np.savetxt(
    csv_path, data, delimiter=",", header=",".join(headers), comments="",
    fmt=["%d", "%.2f", "%d", "%d", "%d", "%.1f", "%d", "%d", "%d", "%d"],
)

workbook = Workbook()
sheet = workbook.active
sheet.title = "churn_data"
sheet.append(headers)
for row in data:
    sheet.append([value.item() if hasattr(value, "item") else value for value in row])

for cell in sheet[1]:
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = PatternFill("solid", fgColor="1D9E75")
sheet.freeze_panes = "A2"
for column in sheet.columns:
    letter = column[0].column_letter
    max_length = max(len(str(cell.value)) for cell in column)
    sheet.column_dimensions[letter].width = min(max(max_length + 2, 13), 30)

workbook.save(xlsx_path)
print(f"Wrote {csv_path.name} and {xlsx_path.name} ({n} rows, churn rate {churn.mean():.0%})")
